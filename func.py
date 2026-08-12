from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os
import time
import tkinter as tk
from tkinter import messagebox, ttk
from datetime import datetime, date
from pandas import DataFrame, read_excel, concat

###############################################################################
###############################################################################

datNuevos = {
    'Fecha': [],
    'Producto': [],
    'Precio': [],
    'Cantidad': [],
    'Tipo': []
}

datosUtil = {
    'Ingresos Totales': [],
    'Egresos Totales': [],
    'Envíos Totales': [],
    'Utilidad Neta': []
}

fechaActual = date.today().strftime("%Y/%m/%d")
fechaActualDos = date.today().strftime("%Y-%m-%d")
horActual = datetime.now().strftime("%H:%M:%S")


###############################################################################
###############################################################################


def formatear_excel(usrActual):

    Excel = load_workbook(f"Reporte-EasyFinance-{usrActual}.xlsx")
    hoja = Excel["Datos"]

    color_fondo = PatternFill(
        fill_type="solid",
        fgColor="81c9fa"
    )   

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = color_fondo
        celda.alignment = Alignment(horizontal="center")

    Excel.save(f"Reporte-EasyFinance-{usrActual}.xlsx")

def login_interfaz(user, passwd):

    Usuario = user.get()
    Contraseña = passwd.get()
    Sesion = (f"{Usuario},{Contraseña}")
    with open("basesdedatos/base_de_datos.txt", "r") as lectarchvo:
        lectura = lectarchvo.read()
    if Sesion in lectura :
        print("Bienvenido al systema")
        from Interfaz import Mostrar_ventana_principal
        Mostrar_ventana_principal()
    else: 
        print("Error")
        tk.messagebox.showerror("Error", "Usuario o contraseña incorrectos. Por favor, inténtelo de nuevo.")
    

def registro_interfaz(user, passwd):

    usr = user.get()
    contrsña = passwd.get()
    regist = (f"{usr},{contrsña}")
    with open("basesdedatos/base_de_datos.txt", "a", encoding="utf-8") as f:
        f.write(regist+"\n")

    tk.messagebox.showinfo("Operacion Completada", "Su cuenta ha sido registrada correctamente. Continue a iniciar sesión.")

def Obtener_fecha_actual():
    meses = {
        1: "Enero",
        2: "Febrero",
        3: "Marzo",
        4: "Abril",
        5: "Mayo",
        6: "Junio",
        7: "Julio",
        8: "Agosto",
        9: "Septiembre",
        10: "Octubre",
        11: "Noviembre",
        12: "Diciembre"
    }

    fechaAct = datetime.now()
    dia = fechaAct.day
    mes = meses[fechaAct.month]
    año = fechaAct.year

    formatFecha = f"{dia} de {mes} de {año}"
    return formatFecha

def registrar_operacion_interfaz(tipoOp, desc, monto, cant, usrAct, util, utilEgr, utilEnv, utilTot, LabelUtlidadTot):

    datNuevos = {
    'Fecha': [],
    'Producto': [],
    'Precio': [],
    'Cantidad': [],
    'Tipo': []
    }

    fecha = date.today().strftime("%Y-%m-%d")
    datNuevos["Fecha"] = fecha

    producto = desc.get()
    datNuevos["Producto"] = producto

    precio = monto.get()
    datNuevos["Precio"] = precio

    cantProductos = cant.get()
    datNuevos["Cantidad"] = cantProductos

    tipo = tipoOp.get()
    datNuevos["Tipo"] = tipo

    if producto == "" or precio == "" or cantProductos == "" or tipo == "" :
        tk.messagebox.showerror("Error", "Ninguno de los campos debe estar vacio para poder registrar una operación")

    else:
        
        try:
            df = read_excel(f"Reporte-EasyFinance-{usrAct}.xlsx", sheet_name="Datos")

        except FileNotFoundError:
            df = DataFrame(columns=["Fecha", "Producto", "Precio", "Cantidad", "Tipo"])

        df = concat([df, DataFrame([datNuevos])], ignore_index=True)
        df.to_excel(f"Reporte-EasyFinance-{usrAct}.xlsx", index=False, sheet_name="Datos")

        formatear_excel(usrAct)

        if tipo == "+(Ingreso)":
            resultUtil = util.get() + (float(precio) * int(cantProductos))
            util.set(f"{resultUtil:.2f}")
        elif tipo == "-(Egreso)":
            resultUtil = utilEgr.get() + (float(precio) * int(cantProductos))
            utilEgr.set(f"{resultUtil:.2f}")
        elif tipo == "-(Envio)":
            resultUtil = utilEnv.get() + (float(precio) * int(cantProductos))
            utilEnv.set(f"{resultUtil:.2f}")
        
        resultado_utilidad_total = util.get() - (utilEgr.get() + utilEnv.get())
        utilTot.set(f"{resultado_utilidad_total:.2f}")

        if utilTot.get() <= 0:
            LabelUtlidadTot.configure(
                fg="#C0503B"
            )
        else:
            LabelUtlidadTot.configure(
                fg="#3FA66B"
            )

        #####################################################################################

        limpiar_widgets(desc)
        limpiar_widgets(monto)
        limpiar_widgets(cant)
        # tipo_operacion.set("")

        #####################################################################################

        tk.messagebox.showinfo("Operacion Registrada", "Su operacion ha sido registrada correctamente en el sistema.")

    return util, utilEgr, utilEnv, utilTot

def limpiar_widgets(widget):
    widget.delete(0, "end")


def definir_color_botones_ingreso(tipOper, regisIngr, registEgr, registEnv):
    if tipOper.get() == "":
        regisIngr.configure(bg="#1D7A6E")
        registEgr.configure(bg="#1D7A6E")
        registEnv.configure(bg="#1D7A6E")

    elif tipOper.get() == "+(Ingreso)":
        regisIngr.configure(bg="#3FA66B")
        registEgr.configure(bg="#1D7A6E")
        registEnv.configure(bg="#1D7A6E")

    elif tipOper.get() == "-(Egreso)":
        registEgr.configure(bg="#3FA66B")
        regisIngr.configure(bg="#1D7A6E")
        registEnv.configure(bg="#1D7A6E")

    elif tipOper.get() == "-(Envio)":
        registEnv.configure(bg="#3FA66B")
        regisIngr.configure(bg="#1D7A6E")
        registEgr.configure(bg="#1D7A6E")

def cargar_tabla_transacciones(Tree, usr):
    for item in Tree.get_children():
        Tree.delete(item)

    try:
        df = read_excel(f"Reporte-EasyFinance-{usr}.xlsx", sheet_name="Datos")

    except FileNotFoundError:
        df = DataFrame(columns=["Fecha", "Producto", "Precio", "Cantidad", "Tipo"])

    if df.empty:
        return

    for index, fila in df.iterrows():

        Tree.insert(
            "",
            "end",
            values=(
                fila["Fecha"],
                fila["Producto"],
                fila["Precio"],
                fila["Cantidad"],
                fila["Tipo"],
        
            ),
        )

def calcular_utilidades_totales(usrAct, varIngr, varEgr, varEnv, varTot, LabelUtlTot):
    ingr = 0
    egr = 0
    env = 0
    total = 0

    try:
        df = read_excel(f"Reporte-EasyFinance-{usrAct}.xlsx", sheet_name="Datos")
    

        for ind, it in enumerate(df["Tipo"]):
            if it == "+(Ingreso)":
                ingr += df["Precio"][ind] * df["Cantidad"][ind]

        for ind, it in enumerate(df["Tipo"]):
            if it == "-(Egreso)":
                egr += df["Precio"][ind] * df["Cantidad"][ind]

        for ind, it in enumerate(df["Tipo"]):
            if it == "-(Envio)":
                env += df["Precio"][ind] * df["Cantidad"][ind]

        total = ingr - (egr + env)

        varIngr.set(f"{ingr:.2f}")
        varEgr.set(f"{egr:.2f}")
        varEnv.set(f"{env:.2f}")
        varTot.set(f"{total:.2f}")

        if varTot.get() <= 0:
            LabelUtlTot.configure(
                fg="#C0503B"
            )
        else:
            LabelUtlTot.configure(
                fg="#3FA66B"
            )

        return varIngr, varEgr, varEnv, varTot

    except:

        varIngr.set(0.00)
        varEgr.set(0.00)
        varEnv.set(0.00)
        varTot.set(0.00)


def actualizar_hora(labelHor):

    actHoraAct = Obtener_fecha_actual()
    labelHor.config(text=f"Nota: La operacion se registrara con la fecha actual. \n ({actHoraAct})")


def logout(varIngr, varEgr, varEnv, varTot):
    with open("basesdedatos/recuerdame.txt", "w", encoding="utf-8") as lectarchv:
        olv = ""
        lectarchv.write(olv)

    varIngr.set("0.00")
    varEgr.set("0.00")
    varEnv.set("0.00")
    varTot.set("0.00")

def exportar_reporte_en_excel(usrAct):
    try:
        tk.messagebox.showinfo("Exportando Reporte", "El reporte de EasyFinance se exportara en su escritorio en formato excel (.xlsx)")
        os.system(f"cp Reporte-EasyFinance-{usrAct}.xlsx ~/Escritorio/Reporte-EasyFinance-{usrAct}.xlsx")
        os.system(f"xdg-open ~/Escritorio/Reporte-EasyFinance-{usrAct}.xlsx")
    except FileNotFoundError:
        os.system(f"cp Reporte-EasyFinance-{usrAct}.xlsx ~/Desktop/Reporte-EasyFinance-{usrAct}.xlsx")
        os.system(f"xdg-open ~/Desktop/Reporte-EasyFinance-{usrAct}.xlsx")


def calcular_Punto_Equilibrio(usrAct):
    try:
        df = read_excel(f"Reporte-EasyFinance-{usrAct}.xlsx", sheet_name="Datos")
    except FileNotFoundError:
        return None

    CF = 0
    ingrTot = 0
    unidVend = 0
    envTot = 0

    for ind, it in enumerate(df["Tipo"]):
        if it == "-(Egreso)":
            CF += df["Precio"][ind] * df["Cantidad"][ind]

    for ind, it in enumerate(df["Tipo"]):
        if it == "+(Ingreso)":
            ingrTot += df["Precio"][ind] * df["Cantidad"][ind]
            unidVend += df["Cantidad"][ind]

    for ind, it in enumerate(df["Tipo"]):
        if it == "-(Envio)":
            envTot += df["Precio"][ind] * df["Cantidad"][ind]

    if unidVend == 0:
        return None

    P = ingrTot / unidVend
    CV = envTot / unidVend

    if P <= CV:
        return {
            "Alcanzable": False,
            "P": P,
            "CV": CV,
            "CF": CF
        }

    xEquil = CF / (P - CV)
    ingrEquil = P * xEquil

    return {
        "Alcanzable": True,
        "X": xEquil,
        "Ingreso": ingrEquil,
        "P": P,
        "CV": CV,
        "CF": CF
    }
